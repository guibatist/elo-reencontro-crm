import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def gerar_apresentacao_elo(dados, arquivos):
    # Inicializa a Apresentação (Formato Widescreen 16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Definição da Paleta de Cores
    COLOR_BG = RGBColor(248, 249, 250)       # Fundo cinza/branco leve
    COLOR_PRIMARY = RGBColor(44, 62, 80)     # Texto primário
    COLOR_ACCENT = RGBColor(142, 185, 214)   # Azul claro calmo
    COLOR_MUTED = RGBColor(127, 140, 141)    # Cinza suave
    COLOR_WHITE = RGBColor(255, 255, 255)    # Branco puro

    def set_slide_background(slide, color):
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = color
        bg_shape.line.fill.background()

    def add_header(slide, title_text, category_text="ELO REENCONTRO"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        p_cat = cat_box.text_frame.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT
        
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        p_title = title_box.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY

    def create_card(slide, left, top, width, height):
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = RGBColor(230, 235, 240)
        card.line.width = Pt(1)

    slide_layout = prs.slide_layouts[6] # Layout em branco

    # ==========================================
    # SLIDE 1: CAPA
    # ==========================================
    slide1 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide1, COLOR_BG)

    left_block = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(4.5), prs.slide_height)
    left_block.fill.solid()
    left_block.fill.fore_color.rgb = COLOR_ACCENT
    left_block.line.fill.background()

    title_box = slide1.shapes.add_textbox(Inches(5.2), Inches(2.2), Inches(7.5), Inches(3.0))
    tf1 = title_box.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "Proposta de Acolhimento\ne Recuperação"
    p1.font.name = "Arial"
    p1.font.size = Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY
    p1.space_after = Pt(24)

    p2 = tf1.add_paragraph()
    p2.text = f"Preparado para: {dados.get('nome_paciente')}\nLocalização: {dados.get('cidade')}"
    p2.font.name = "Arial"
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLOR_MUTED

    # Logo Elo Reencontro Fixo
    if os.path.exists("static/logo.png"):
        slide1.shapes.add_picture("static/logo.png", Inches(0.8), Inches(6.2), width=Inches(2.9), height=Inches(0.6))
    else:
        logo1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.2), Inches(2.9), Inches(0.6))
        logo1.fill.solid()
        logo1.fill.fore_color.rgb = COLOR_WHITE
        logo1.text_frame.text = "ELO REENCONTRO"

    # Logo da Clínica Dinâmico
    if arquivos.get('logo_clinica'):
        slide1.shapes.add_picture(arquivos.get('logo_clinica'), Inches(5.2), Inches(6.2), width=Inches(2.9), height=Inches(0.6))
    else:
        logo2 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.2), Inches(6.2), Inches(2.9), Inches(0.6))
        logo2.fill.solid()
        logo2.fill.fore_color.rgb = COLOR_WHITE
        logo2.text_frame.text = "[ LOGO DA CLÍNICA ]"

    # ==========================================
    # SLIDE 2: A CIÊNCIA DA MUDANÇA
    # ==========================================
    slide2 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide2, COLOR_BG)
    add_header(slide2, "Entendendo o Mecanismo do Vício", "A Ciência da Mudança")

    text_box2 = slide2.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(6.0), Inches(4.5))
    tf2 = text_box2.text_frame
    tf2.word_wrap = True
    p_neuro = tf2.paragraphs[0]
    p_neuro.text = "A neurobiologia da dependência química não é uma questão de escolha, mas de alteração estrutural no cérebro."
    p_neuro.font.size = Pt(16)
    p_neuro.font.bold = True
    p_neuro.font.color.rgb = COLOR_PRIMARY
    p_neuro.space_after = Pt(14)

    p_neuro_sub = tf2.add_paragraph()
    p_neuro_sub.text = "O uso contínuo de substâncias sequestra o sistema de recompensa cerebral, superestimulando a liberação de dopamina. Com o tempo, o cérebro perde a capacidade de sentir prazer em atividades comuns, gerando a compulsão física e psicológica."
    p_neuro_sub.font.size = Pt(13)
    p_neuro_sub.font.color.rgb = COLOR_MUTED

    pillars = [
        ("Desintoxicação e Reset", "Interrupção do ciclo químico para limpar os receptores cerebrais com suporte médico."),
        ("Estimulação da Neuroplasticidade", "Terapias cognitivas que ajudam o cérebro a criar novos caminhos e hábitos saudáveis."),
        ("Ambiente Controlado", "Isolamento de gatilhos cotidianos, permitindo que a química cerebral retorne ao seu equilíbrio natural.")
    ]
    for i, (title, desc) in enumerate(pillars):
        top_pos = Inches(2.0 + (i * 1.5))
        create_card(slide2, Inches(7.3), top_pos, Inches(5.2), Inches(1.3))
        tbox = slide2.shapes.add_textbox(Inches(7.5), top_pos + Inches(0.1), Inches(4.8), Inches(1.1))
        tf = tbox.text_frame
        tf.word_wrap = True
        p_title = tf.paragraphs[0]
        p_title.text = f"• {title}"
        p_title.font.bold = True
        p_title.font.size = Pt(14)
        p_title.font.color.rgb = COLOR_PRIMARY
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = COLOR_MUTED

    # ==========================================
    # SLIDE 3: O PROTAGONISMO DO PACIENTE
    # ==========================================
    slide3 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide3, COLOR_BG)
    add_header(slide3, "O Despertar do Paciente", "O Protagonismo do Paciente")

    patient_pillars = [
        ("Aceitação Pragmática", "Reconhecer a vulnerabilidade diante da substância é o primeiro passo para a libertação definitiva."),
        ("Autorresponsabilidade", "O tratamento transfere gradativamente o controle de volta ao paciente, transformando-o em agente da cura."),
        ("Reconstrução da Identidade", "Mais do que parar de usar, o foco está em redescobrir valores, propósitos e talentos esquecidos.")
    ]
    for i, (title, desc) in enumerate(patient_pillars):
        left_pos = Inches(0.8 + (i * 4.0))
        create_card(slide3, left_pos, Inches(2.2), Inches(3.7), Inches(4.2))
        accent_bar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, Inches(2.2), Inches(3.7), Inches(0.1))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = COLOR_ACCENT
        accent_bar.line.fill.background()
        
        tbox = slide3.shapes.add_textbox(left_pos + Inches(0.2), Inches(2.5), Inches(3.3), Inches(3.6))
        tf = tbox.text_frame
        tf.word_wrap = True
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.bold = True
        p_title.font.size = Pt(16)
        p_title.font.color.rgb = COLOR_PRIMARY
        p_title.space_after = Pt(12)
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = COLOR_MUTED

    # ==========================================
    # SLIDE 4: O APOIO FAMILIAR
    # ==========================================
    slide4 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide4, COLOR_BG)
    add_header(slide4, "A Família como Porto Seguro", "O Apoio Familiar")

    left_box4 = slide4.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(5.0), Inches(4.0))
    tf4_left = left_box4.text_frame
    tf4_left.word_wrap = True
    p_fam_intro = tf4_left.paragraphs[0]
    p_fam_intro.text = "A reabilitação é um processo sistêmico.\nQuando a família muda, o paciente encontra solo fértil para florescer."
    p_fam_intro.font.size = Pt(20)
    p_fam_intro.font.bold = True
    p_fam_intro.font.color.rgb = COLOR_ACCENT
    p_fam_intro.space_after = Pt(14)

    family_steps = [
        ("Limites Saudáveis", "Aprender a dizer não e proteger a integridade emocional do lar, transformando o amor que aceita em um amor que protege."),
        ("Fim da Codependência", "Cessar os comportamentos que inconscientemente protegem ou escondem a dependência do paciente, quebrando ciclos nocivos."),
        ("Participação Ativa", "Engajamento em reuniões familiares, grupos de apoio dedicados e psicoeducação promovidos pela instituição.")
    ]
    for i, (title, desc) in enumerate(family_steps):
        top_pos = Inches(2.2 + (i * 1.5))
        sq = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.5), top_pos + Inches(0.05), Inches(0.15), Inches(0.15))
        sq.fill.solid()
        sq.fill.fore_color.rgb = COLOR_ACCENT
        sq.line.fill.background()
        tbox = slide4.shapes.add_textbox(Inches(6.9), top_pos, Inches(5.6), Inches(1.3))
        tf = tbox.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.bold = True
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = COLOR_PRIMARY
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_MUTED

    # ==========================================
    # SLIDE 5: ESTRUTURA DE EXCELÊNCIA
    # ==========================================
    slide5 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide5, COLOR_BG)
    add_header(slide5, f"Unidade: {dados.get('nome_clinica')}", "Estrutura de Excelência")

    features_box = slide5.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(5.5), Inches(4.5))
    tf5 = features_box.text_frame
    tf5.word_wrap = True
    features = [
        ("Equipe Multidisciplinar", "Médicos psiquiatras, psicólogos especialistas em dependência, terapeutas e enfermagem 24 horas."),
        ("Hotelaria e Conforto", "Acomodações acolhedoras projetadas para garantir repouso, dignidade e bem-estar em áreas verdes."),
        ("Segurança Monitorada", "Protocolos rigorosos de acolhimento e monitoramento para assegurar a total integridade física.")
    ]
    for i, (title, desc) in enumerate(features):
        p_t = tf5.paragraphs[0] if i == 0 else tf5.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = COLOR_PRIMARY
        if i > 0: p_t.space_before = Pt(12)
        p_d = tf5.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_MUTED

    # Foto Estrutura Dinâmica
    if arquivos.get('foto_estrutura'):
        slide5.shapes.add_picture(arquivos.get('foto_estrutura'), Inches(7.0), Inches(2.2), width=Inches(5.5), height=Inches(4.2))
    else:
        photo_box = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.0), Inches(2.2), Inches(5.5), Inches(4.2))
        photo_box.fill.solid()
        photo_box.fill.fore_color.rgb = RGBColor(235, 240, 245)
        photo_box.text_frame.text = "[ SEM FOTO DA ESTRUTURA ]"

    # ==========================================
    # SLIDE 6: A CONEXÃO ELO REENCONTRO
    # ==========================================
    slide6 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide6, COLOR_BG)
    add_header(slide6, "Por que esta escolha?", "A Conexão Elo Reencontro")

    connections = [
        ("Curadoria Técnica e Humanizada", "Avaliamos os índices reais de recuperação da unidade, a linha terapêutica e a sinergia com o perfil do paciente. Esta clínica foi selecionada de forma cirúrgica para as necessidades da família."),
        ("O Acompanhamento Contínuo", "Atuamos como o elo seguro de comunicação entre a clínica e a família. Monitoramos os relatórios de evolução, fornecemos suporte aos decisores e ajudamos no planejamento da alta.")
    ]
    for i, (title, desc) in enumerate(connections):
        top_pos = Inches(2.2 + (i * 2.3))
        create_card(slide6, Inches(0.8), top_pos, Inches(11.733), Inches(1.9))
        bar = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top_pos, Inches(0.1), Inches(1.9))
        bar.fill.solid()
        bar.fill.fore_color.rgb = COLOR_ACCENT
        bar.line.fill.background()
        
        tbox = slide6.shapes.add_textbox(Inches(1.2), top_pos + Inches(0.15), Inches(11.0), Inches(1.6))
        tf = tbox.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.bold = True
        p_t.font.size = Pt(16)
        p_t.font.color.rgb = COLOR_PRIMARY
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_MUTED

    # ==========================================
    # SLIDE 7: PRÓXIMOS PASSOS (DADOS FIXADOS)
    # ==========================================
    slide7 = prs.slides.add_slide(slide_layout)
    set_slide_background(slide7, COLOR_PRIMARY)

    title_box7 = slide7.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(1.0))
    p7_title = title_box7.text_frame.paragraphs[0]
    p7_title.text = "O Reencontro Começa Agora"
    p7_title.alignment = PP_ALIGN.CENTER
    p7_title.font.size = Pt(36)
    p7_title.font.bold = True
    p7_title.font.color.rgb = COLOR_WHITE

    phrase_box = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.5), Inches(3.6), Inches(8.333), Inches(0.9))
    phrase_box.fill.solid()
    phrase_box.fill.fore_color.rgb = RGBColor(52, 73, 94)
    phrase_box.line.fill.background()
    p_phrase = phrase_box.text_frame.paragraphs[0]
    p_phrase.text = "“Não fazemos apenas internações, reconectamos vidas.”"
    p_phrase.alignment = PP_ALIGN.CENTER
    p_phrase.font.size = Pt(16)
    p_phrase.font.italic = True
    p_phrase.font.bold = True
    p_phrase.font.color.rgb = COLOR_WHITE

    contact_box = slide7.shapes.add_textbox(Inches(1.0), Inches(5.5), Inches(11.333), Inches(0.8))
    p7_contact = contact_box.text_frame.paragraphs[0]
    p7_contact.text = "Contato Direto: (11) 96208-4852   |   E-mail: contato@eloreencontro.com.br"
    p7_contact.alignment = PP_ALIGN.CENTER
    p7_contact.font.size = Pt(14)
    p7_contact.font.color.rgb = COLOR_WHITE

    # ==========================================
    # SLIDES EXTRAS: 5 FOLHAS DE FOTOS QUASE FULLSCREEN
    # ==========================================
    for idx in range(1, 6):
        chave_foto = f'foto_extra_{idx}'
        if arquivos.get(chave_foto):
            slide_extra = prs.slides.add_slide(slide_layout)
            set_slide_background(slide_extra, COLOR_BG)
            add_header(slide_extra, f"Galeria de Fotos da Unidade - Imagem {idx}", "Infraestrutura Completa")
            
            # Adiciona a imagem de forma generosa centralizada na página
            slide_extra.shapes.add_picture(
                arquivos.get(chave_foto), 
                Inches(1.0), Inches(1.8),  # Margens esquerda e superior
                width=Inches(11.333), height=Inches(5.0) # Tamanho estendido
            )

    return prs